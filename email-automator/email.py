#!/usr/bin/env python3
"""Send personalized emails from BioMod Gmail accounts using a CSV and a Word template."""

from __future__ import annotations

import argparse
import base64
import csv
import logging
import re
import sys
import time
import traceback
from datetime import datetime, timezone
from html import escape as html_escape
from pathlib import Path

# This file is named email.py, which would otherwise hide Python's built-in email package.
def _import_stdlib_email_mime():
    script_dir = str(Path(__file__).resolve().parent)
    saved_path = sys.path[:]
    sys.path = [entry for entry in sys.path if entry not in ("", script_dir)]
    try:
        from email.mime.multipart import MIMEMultipart
        from email.mime.text import MIMEText

        return MIMEMultipart, MIMEText
    finally:
        sys.path = saved_path


MIMEMultipart, MIMEText = _import_stdlib_email_mime()

SCRIPT_DIR = Path(__file__).resolve().parent
ALLOWED_SENDERS = (
    "ubcbiomod@gmail.com",
    "wetlab.ubcbiomod@gmail.com",
)
GMAIL_SCOPES = (
    "https://www.googleapis.com/auth/gmail.send",
    "https://www.googleapis.com/auth/userinfo.email",
    "openid",
)
PLACEHOLDER_RE = re.compile(r"\{\{\s*([^}]+?)\s*\}\}")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
SENT_LOG_FIELDS = (
    "timestamp",
    "campaign",
    "sender",
    "row",
    "email",
    "status",
    "message",
)


class UserFacingError(Exception):
    """An error that should be shown in plain English, without a traceback."""


def configure_stdio() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def setup_debug_log(path: Path) -> None:
    logging.basicConfig(
        filename=path,
        filemode="a",
        level=logging.DEBUG,
        format="%(asctime)s %(levelname)s %(message)s",
        encoding="utf-8",
    )


def say(message: str) -> None:
    print(message)
    logging.info(message)


def warn(message: str) -> None:
    print(message)
    logging.warning(message)


def fail(message: str) -> None:
    logging.error(message)
    raise UserFacingError(message)


def token_path_for(account: str) -> Path:
    safe = account.strip().lower().replace("@", "_at_").replace(".", "_")
    return SCRIPT_DIR / f"token_{safe}.json"


def normalize_header(value: str) -> str:
    cleaned = value.strip().lower().replace("_", " ")
    return re.sub(r"\s+", " ", cleaned)


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Send personalized emails from a CSV and a Word (.docx) template."
    )
    parser.add_argument(
        "--config",
        default=str(SCRIPT_DIR / "config.yaml"),
        help="Path to config.yaml",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Preview emails without sending",
    )
    parser.add_argument(
        "--create-examples",
        action="store_true",
        help="Create a sample recipients.csv and template.docx in this folder",
    )
    return parser.parse_args()


def load_config(path: Path) -> dict:
    try:
        import yaml
    except ImportError:
        fail(
            "The PyYAML package is missing. Open a terminal in this folder and run:\n"
            "  python -m pip install -r requirements.txt"
        )

    if not path.exists():
        fail(
            f"Can't find the settings file at {path}. "
            "Keep config.yaml in the same folder as email.py, or pass --config with the correct path."
        )

    try:
        with path.open(encoding="utf-8") as handle:
            data = yaml.safe_load(handle) or {}
    except Exception as exc:
        logging.exception("Failed to parse config")
        fail(
            f"Could not read {path.name}. Make sure it is valid YAML "
            f"(check quotes and indentation). Technical detail saved in debug.log ({exc})."
        )

    if not isinstance(data, dict):
        fail(f"{path.name} should contain settings like sender_account and csv_file, not a list.")
    return data


def required_setting(config: dict, key: str) -> str:
    value = config.get(key)
    if value is None or str(value).strip() == "":
        fail(
            f"config.yaml is missing '{key}'. Open config.yaml in Notepad and fill it in. "
            "See README.md for what each setting means."
        )
    return str(value).strip()


def resolve_path(raw: str) -> Path:
    path = Path(raw)
    if not path.is_absolute():
        path = SCRIPT_DIR / path
    return path.resolve()


def create_example_files() -> None:
    csv_path = SCRIPT_DIR / "recipients.csv"
    template_path = SCRIPT_DIR / "template.docx"
    example_csv = SCRIPT_DIR / "recipients.example.csv"

    if not csv_path.exists():
        if example_csv.exists():
            csv_path.write_text(example_csv.read_text(encoding="utf-8"), encoding="utf-8")
        else:
            csv_path.write_text(
                "email,first_name,team_name\n"
                "jane@example.com,Jane,UBC iGEM\n"
                "alex@example.com,Alex,Wetlab\n",
                encoding="utf-8",
            )
        say(f"Created sample recipient list: {csv_path.name}")
    else:
        say(f"{csv_path.name} already exists — left it unchanged.")

    if not template_path.exists():
        try:
            from docx import Document
        except ImportError:
            fail(
                "python-docx is missing, so a sample Word file could not be created. "
                "Run: python -m pip install -r requirements.txt"
            )
        document = Document()
        document.add_paragraph("Hi {{first_name}},")
        document.add_paragraph("")
        document.add_paragraph(
            "This is a sample email for {{team_name}}. Replace this text with your real message. "
            "Names inside double curly braces are filled in from your spreadsheet."
        )
        document.add_paragraph("")
        document.add_paragraph("Thanks,")
        document.add_paragraph("UBC BioMod")
        document.save(template_path)
        say(f"Created sample Word template: {template_path.name}")
    else:
        say(f"{template_path.name} already exists — left it unchanged.")

    say("You can now edit those two files, then run the sender as usual.")


def read_recipients(csv_path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if not csv_path.exists():
        fail(
            f"Can't find the CSV file at {csv_path}. "
            "Check that the file exists and that csv_file in config.yaml matches its name."
        )
    if csv_path.suffix.lower() != ".csv":
        fail(
            f"The recipient list must be a .csv file. You pointed csv_file at {csv_path.name}."
        )

    try:
        with csv_path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames:
                fail(
                    f"{csv_path.name} has no header row. The first row should be column names "
                    "such as email,first_name,team_name."
                )
            original_headers = [h for h in reader.fieldnames if h is not None]
            normalized_headers = [normalize_header(h) for h in original_headers]
            if len(set(normalized_headers)) != len(normalized_headers):
                fail(
                    f"{csv_path.name} has duplicate column names (ignoring capitalization). "
                    f"Found columns: {', '.join(original_headers)}."
                )
            header_map = {
                normalize_header(original): original for original in original_headers
            }
            if "email" not in header_map:
                fail(
                    f"Missing 'email' column in CSV — found columns: {', '.join(original_headers)}. "
                    "Please add an 'email' column."
                )

            rows: list[dict[str, str]] = []
            for index, raw_row in enumerate(reader, start=2):
                row = {
                    normalize_header(key): (value or "").strip()
                    for key, value in raw_row.items()
                    if key is not None
                }
                row["_row_number"] = str(index)
                rows.append(row)
    except UserFacingError:
        raise
    except Exception as exc:
        logging.exception("Failed to read CSV")
        fail(
            f"Could not read {csv_path.name}. Make sure it is a CSV saved from Excel "
            f"(File → Save As → CSV UTF-8). Technical detail saved in debug.log ({exc})."
        )

    if not rows:
        fail(f"{csv_path.name} has a header row but no recipients. Add at least one person.")
    return original_headers, rows


def extract_template_text(template_path: Path) -> str:
    if not template_path.exists():
        fail(
            f"Can't find the Word template at {template_path}. "
            "Check that the file exists and that template_file in config.yaml matches its name."
        )
    if template_path.suffix.lower() != ".docx":
        fail(
            f"The email template must be a Word .docx file (not .doc, .pdf, or .txt). "
            f"You pointed template_file at {template_path.name}."
        )

    try:
        from docx import Document
    except ImportError:
        fail(
            "The python-docx package is missing. Open a terminal in this folder and run:\n"
            "  python -m pip install -r requirements.txt"
        )

    try:
        document = Document(template_path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for table_row in table.rows:
                for cell in table_row.cells:
                    paragraphs.extend(cell.paragraphs and [p.text for p in cell.paragraphs] or [])
    except Exception as exc:
        logging.exception("Failed to read Word template")
        fail(
            f"Could not open {template_path.name}. Save it again from Microsoft Word as "
            f"a .docx file. Technical detail saved in debug.log ({exc})."
        )

    text = "\n".join(paragraphs).strip()
    if not text:
        fail(
            f"{template_path.name} appears to be empty. Type your email in Word, using "
            "placeholders like {{first_name}}, then save and try again."
        )
    return text


def find_placeholders(*texts: str) -> list[tuple[str, str]]:
    found: list[tuple[str, str]] = []
    seen: set[str] = set()
    for text in texts:
        for match in PLACEHOLDER_RE.finditer(text):
            original = match.group(1).strip()
            name = normalize_header(original)
            if name and name not in seen:
                seen.add(name)
                found.append((name, original))
    return found


def substitute(text: str, row: dict[str, str]) -> str:
    def replace(match: re.Match[str]) -> str:
        key = normalize_header(match.group(1))
        return row.get(key, "")

    return PLACEHOLDER_RE.sub(replace, text)


def validate_placeholders(placeholders: list[tuple[str, str]], csv_headers: list[str]) -> None:
    available = {normalize_header(header) for header in csv_headers}
    missing = [original for name, original in placeholders if name not in available]
    if missing:
        pretty = ", ".join(f"{{{{{name}}}}}" for name in missing)
        fail(
            f"Template uses {pretty} but CSV has no matching column — "
            f"found columns: {', '.join(csv_headers)}. "
            "Check your CSV headers or template placeholders."
        )


def is_plausible_email(value: str) -> bool:
    return bool(EMAIL_RE.match(value))


def load_sent_log(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    try:
        with path.open(newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle)
            return [dict(row) for row in reader]
    except Exception as exc:
        logging.exception("Failed to read sent log")
        fail(
            f"Could not read {path.name}. If this file looks damaged, rename it "
            f"(for example to sent_log.bak.csv) and run again. Technical detail saved in debug.log ({exc})."
        )
    return []


def append_sent_log(path: Path, record: dict[str, str]) -> None:
    new_file = not path.exists()
    with path.open("a", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=SENT_LOG_FIELDS)
        if new_file:
            writer.writeheader()
        writer.writerow({field: record.get(field, "") for field in SENT_LOG_FIELDS})


def already_sent_emails(log_rows: list[dict[str, str]], campaign: str, sender: str) -> set[str]:
    sent: set[str] = set()
    for row in log_rows:
        if (
            row.get("campaign", "").strip() == campaign
            and row.get("sender", "").strip().lower() == sender
            and row.get("status", "").strip().lower() == "sent"
        ):
            sent.add(row.get("email", "").strip().lower())
    return sent


def sent_today_count(log_rows: list[dict[str, str]], sender: str, today: datetime) -> int:
    count = 0
    today_str = today.date().isoformat()
    for row in log_rows:
        if row.get("sender", "").strip().lower() != sender:
            continue
        if row.get("status", "").strip().lower() != "sent":
            continue
        timestamp = row.get("timestamp", "")
        if timestamp.startswith(today_str):
            count += 1
    return count


def preview_email(index: int, recipient: dict[str, str], subject: str, body: str) -> None:
    say("")
    say(f"----- Preview {index} -----")
    say(f"To: {recipient['email']}")
    say(f"Subject: {subject}")
    say("")
    say(body)
    say("----- End preview -----")


def confirm_send(count: int, sender: str) -> bool:
    say("")
    prompt = f"You are about to send {count} emails from {sender}. Continue? [y/n] "
    try:
        answer = input(prompt).strip().lower()
    except EOFError:
        fail("No confirmation was provided, so nothing was sent.")
    return answer in {"y", "yes"}


def import_google() -> tuple:
    try:
        from google.auth.transport.requests import Request
        from google.oauth2.credentials import Credentials
        from google_auth_oauthlib.flow import InstalledAppFlow
        from googleapiclient.discovery import build
        from googleapiclient.errors import HttpError
    except ImportError:
        fail(
            "Google Gmail libraries are missing. Open a terminal in this folder and run:\n"
            "  python -m pip install -r requirements.txt"
        )
    return Request, Credentials, InstalledAppFlow, build, HttpError


def authenticate(sender: str, credentials_path: Path, token_path: Path):
    Request, Credentials, InstalledAppFlow, build, _HttpError = import_google()

    if not credentials_path.exists():
        fail(
            f"Can't find Google login file credentials.json at {credentials_path}. "
            "A one-time Google Cloud setup is required. Follow the 'Connect Gmail' section in README.md, "
            "then put the downloaded credentials.json in this folder."
        )

    creds = None
    if token_path.exists():
        try:
            creds = Credentials.from_authorized_user_file(str(token_path), GMAIL_SCOPES)
        except Exception:
            logging.exception("Could not load saved token")
            warn(
                f"The saved login file {token_path.name} could not be read, so you will need to log in again."
            )
            creds = None

    try:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        if not creds or not creds.valid:
            say("A browser window will open so you can log in to Gmail. Use the account listed in config.yaml.")
            flow = InstalledAppFlow.from_client_secrets_file(str(credentials_path), GMAIL_SCOPES)
            creds = flow.run_local_server(port=0)
        token_path.write_text(creds.to_json(), encoding="utf-8")
    except UserFacingError:
        raise
    except Exception:
        logging.exception("OAuth failed")
        fail(
            f"Could not log in to {sender} — your saved login may have expired or the Google setup is incomplete. "
            "Run the script again and log in when the browser opens. If a browser never appears, "
            "follow the 'Connect Gmail' steps in README.md. Technical detail is in debug.log."
        )

    try:
        oauth = build("oauth2", "v2", credentials=creds, cache_discovery=False)
        profile = oauth.userinfo().get().execute()
        logged_in = (profile.get("email") or "").strip().lower()
    except Exception:
        logging.exception("Could not read logged-in email")
        fail(
            f"Logged in, but could not confirm which Gmail account was used. "
            f"Delete {token_path.name} and run again, then choose {sender} in the browser."
        )

    if logged_in != sender:
        fail(
            f"You logged in as {logged_in}, but config.yaml says to send from {sender}. "
            f"Delete {token_path.name} in this folder, then run again and choose {sender} in the browser."
        )

    gmail = build("gmail", "v1", credentials=creds, cache_discovery=False)
    say(f"Logged in as {sender}.")
    return gmail


def build_message(sender: str, to_address: str, subject: str, body: str) -> dict:
    message = MIMEMultipart("alternative")
    message["To"] = to_address
    message["From"] = sender
    message["Subject"] = subject
    message.attach(MIMEText(body, "plain", "utf-8"))
    html_body = "".join(f"<p>{html_escape(line) if line else '&nbsp;'}</p>" for line in body.split("\n"))
    message.attach(MIMEText(html_body, "html", "utf-8"))
    raw = base64.urlsafe_b64encode(message.as_bytes()).decode("ascii")
    return {"raw": raw}


def describe_send_error(exc: Exception, HttpError) -> tuple[str, bool, bool]:
    """Return (human message, daily quota hit, short-window rate limit)."""
    if HttpError is not None and isinstance(exc, HttpError):
        status = getattr(exc.resp, "status", None)
        content = ""
        try:
            content = exc.content.decode("utf-8", errors="replace") if exc.content else ""
        except Exception:
            content = str(exc)
        lowered = content.lower()
        daily_markers = ("dailylimitexceeded", "daily limit", "quotaexceeded")
        rate_markers = ("user_rate_limit_exceeded", "ratelimitexceeded", "rate limit")
        if status in (429, 403) and any(marker in lowered for marker in daily_markers):
            return "Gmail daily send limit reached", True, False
        if status in (429, 403) and any(marker in lowered for marker in rate_markers):
            return "Gmail asked us to slow down (rate limit)", False, True
        if status == 400:
            return "Gmail rejected the address or message as invalid", False, False
        return f"Gmail returned an error (code {status})", False, False
    return "network or unexpected Gmail error", False, False


def send_one(gmail, sender: str, to_address: str, subject: str, body: str) -> None:
    gmail.users().messages().send(userId="me", body=build_message(sender, to_address, subject, body)).execute()


def prepare_queue(
    rows: list[dict[str, str]],
    already_sent: set[str],
    sent_log_path: Path,
    campaign: str,
    sender: str,
) -> list[dict[str, str]]:
    queue: list[dict[str, str]] = []
    for row in rows:
        row_number = row["_row_number"]
        address = row.get("email", "").strip()
        if not address:
            message = f"Row {row_number} skipped: 'email' field is blank."
            warn(message)
            append_sent_log(
                sent_log_path,
                {
                    "timestamp": utc_now().isoformat(timespec="seconds"),
                    "campaign": campaign,
                    "sender": sender,
                    "row": row_number,
                    "email": "",
                    "status": "skipped",
                    "message": message,
                },
            )
            continue
        if not is_plausible_email(address):
            message = (
                f"Row {row_number} skipped: '{address}' does not look like a valid email address."
            )
            warn(message)
            append_sent_log(
                sent_log_path,
                {
                    "timestamp": utc_now().isoformat(timespec="seconds"),
                    "campaign": campaign,
                    "sender": sender,
                    "row": row_number,
                    "email": address,
                    "status": "skipped",
                    "message": message,
                },
            )
            continue
        if address.lower() in already_sent:
            message = (
                f"Row {row_number} skipped: {address} already has a successful send "
                f"logged for campaign '{campaign}'."
            )
            say(message)
            continue
        queue.append(row)
    return queue


def run(args: argparse.Namespace) -> int:
    debug_path = SCRIPT_DIR / "debug.log"
    setup_debug_log(debug_path)
    logging.info("Starting email sender")

    if args.create_examples:
        create_example_files()
        return 0

    config = load_config(Path(args.config))
    sender = required_setting(config, "sender_account").lower()
    # if sender not in ALLOWED_SENDERS:
    #     fail(
    #         f"sender_account in config.yaml is '{sender}', which is not allowed. "
    #         f"Use one of: {', '.join(ALLOWED_SENDERS)}."
    #     )

    csv_path = resolve_path(required_setting(config, "csv_file"))
    template_path = resolve_path(required_setting(config, "template_file"))
    subject_template = required_setting(config, "subject")
    campaign = str(config.get("campaign") or "default").strip() or "default"
    preview_count = int(config.get("preview_count") or 2)
    delay_seconds = float(config.get("delay_seconds") or 1.5)
    daily_limit = int(config.get("daily_limit") or 500)
    dry_run = bool(args.dry_run or config.get("dry_run"))
    credentials_path = SCRIPT_DIR / "credentials.json"
    sent_log_path = SCRIPT_DIR / "sent_log.csv"

    csv_headers, rows = read_recipients(csv_path)
    template_text = extract_template_text(template_path)
    placeholders = find_placeholders(template_text, subject_template)
    validate_placeholders(placeholders, csv_headers)

    log_rows = load_sent_log(sent_log_path)
    already_sent = already_sent_emails(log_rows, campaign, sender)
    used_today = sent_today_count(log_rows, sender, utc_now())
    remaining_today = max(0, daily_limit - used_today)

    queue = prepare_queue(rows, already_sent, sent_log_path, campaign, sender)
    if remaining_today <= 0:
        fail(
            f"Daily send limit ({daily_limit}) reached for {sender}. "
            "Remaining recipients will need to be sent tomorrow or from the other account."
        )
    if len(queue) > remaining_today:
        warn(
            f"{sender} can send {remaining_today} more email(s) today (limit {daily_limit}). "
            f"Only the first {remaining_today} of {len(queue)} remaining recipients will be sent."
        )
        queue = queue[:remaining_today]

    if not queue:
        say(
            "There is nobody left to email. Either the CSV is empty after skipping invalid rows, "
            "or everyone was already marked as sent in sent_log.csv for this campaign."
        )
        return 0

    say(f"Ready to send from: {sender}")
    say(f"Campaign: {campaign}")
    say(f"Recipients in this run: {len(queue)}")
    say(f"Already sent today from this account (from the log): {used_today} of {daily_limit}")

    samples = queue[: max(1, preview_count)]
    for index, recipient in enumerate(samples, start=1):
        preview_email(
            index,
            recipient,
            substitute(subject_template, recipient),
            substitute(template_text, recipient),
        )

    if dry_run:
        say("")
        say("Dry run only — no emails were sent. Set dry_run: false in config.yaml when you are ready.")
        return 0

    if not confirm_send(len(queue), sender):
        say("Cancelled. No emails were sent.")
        return 0

    gmail = authenticate(sender, credentials_path, token_path_for(sender))
    _, _, _, _, HttpError = import_google()

    sent_ok = 0
    failed = 0
    for index, recipient in enumerate(queue, start=1):
        address = recipient["email"]
        row_number = recipient["_row_number"]
        subject = substitute(subject_template, recipient)
        body = substitute(template_text, recipient)
        try:
            send_one(gmail, sender, address, subject, body)
            message = f"Sent to {address} (row {row_number})."
            say(f"[{index}/{len(queue)}] {message}")
            append_sent_log(
                sent_log_path,
                {
                    "timestamp": utc_now().isoformat(timespec="seconds"),
                    "campaign": campaign,
                    "sender": sender,
                    "row": row_number,
                    "email": address,
                    "status": "sent",
                    "message": message,
                },
            )
            sent_ok += 1
        except Exception as exc:
            logging.exception("Send failed for %s row %s", address, row_number)
            reason, quota_hit, rate_limited = describe_send_error(exc, HttpError)
            if rate_limited:
                warn(
                    f"Failed to send to {address} (row {row_number}): {reason}. "
                    "Waiting 30 seconds, then continuing with the rest of the list."
                )
                append_sent_log(
                    sent_log_path,
                    {
                        "timestamp": utc_now().isoformat(timespec="seconds"),
                        "campaign": campaign,
                        "sender": sender,
                        "row": row_number,
                        "email": address,
                        "status": "failed",
                        "message": reason,
                    },
                )
                failed += 1
                time.sleep(30)
                continue
            if quota_hit:
                message = (
                    f"Daily send limit ({daily_limit}) reached for {sender}. "
                    "Remaining recipients will need to be sent tomorrow or from the other account."
                )
                warn(
                    f"Failed to send to {address} (row {row_number}): {reason}. {message}"
                )
                append_sent_log(
                    sent_log_path,
                    {
                        "timestamp": utc_now().isoformat(timespec="seconds"),
                        "campaign": campaign,
                        "sender": sender,
                        "row": row_number,
                        "email": address,
                        "status": "failed",
                        "message": reason,
                    },
                )
                failed += 1
                say(message)
                break
            message = f"Failed to send to {address} (row {row_number}): {reason}."
            warn(message)
            append_sent_log(
                sent_log_path,
                {
                    "timestamp": utc_now().isoformat(timespec="seconds"),
                    "campaign": campaign,
                    "sender": sender,
                    "row": row_number,
                    "email": address,
                    "status": "failed",
                    "message": reason,
                },
            )
            failed += 1
        if index < len(queue):
            time.sleep(max(0.0, delay_seconds))

    say("")
    say(f"Finished. Sent: {sent_ok}. Failed or skipped during send: {failed}.")
    say(f"A full record is in {sent_log_path.name}. Technical details (if any) are in {debug_path.name}.")
    return 0 if failed == 0 else 1


def main() -> int:
    configure_stdio()
    args = parse_args()
    try:
        return run(args)
    except UserFacingError as exc:
        print()
        print(str(exc))
        print()
        print("Nothing else was changed. If you need help, share debug.log and sent_log.csv with a teammate.")
        return 1
    except KeyboardInterrupt:
        print()
        print("Stopped. Emails already sent stay sent; check sent_log.csv before running again.")
        return 1
    except Exception:
        logging.exception("Unhandled error")
        print()
        print(
            "Something unexpected went wrong. The technical details were saved in debug.log "
            "in this folder — you can send that file to a teammate. "
            f"({traceback.format_exc().splitlines()[-1]})"
        )
        return 1


if __name__ == "__main__":
    sys.exit(main())
